"""Read wording edited in PowerPoint back out of the deck.

The deck is generated from diecast_content.js, so text typed into the .pptx is
lost on the next build. This reads the deck, works out which content field each
text box came from, and prints (or applies) the differences.

    python tools/training-deck/pull_pptx_edits.py            # show what changed
    python tools/training-deck/pull_pptx_edits.py --apply    # write it into the content file

Two things make this less obvious than it looks:
  * slides can be REORDERED in PowerPoint, so position identifies nothing --
    slides are matched by title, and the content file is reordered to match.
  * marker discs and zone letters are bold text boxes ('**1**', '**A**'), so
    they must be dropped before fields are lined up, or everything shifts.
Bold runs come back as **bold**, the way the content file marks screen labels.
"""
import json
import re
import subprocess
import sys
import pathlib
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

HERE = pathlib.Path(__file__).resolve().parent
SHOTS = HERE.parent.parent / 'docs/training/diecast/shots'
DECK = HERE.parent.parent / 'docs/training/diecast/MPP_DieCast_Training.pptx'
CONTENT = HERE / 'diecast_content.js'


def content_slides():
    out = subprocess.run(
        ['node', '-e', f'process.stdout.write(JSON.stringify(require({json.dumps(str(CONTENT))}).slides))'],
        capture_output=True, text=True, check=True)
    return json.loads(out.stdout)


SMART = {'‘': "'", '’': "'", '“': '"', '”': '"',
         '–': '-', '—': '--', '…': '...', ' ': ' ', '·': '-'}


def ascii_fold(t):
    """PowerPoint inserts curly quotes and dashes; content must stay ASCII
    (CLAUDE.md: sqlcmd turns non-ASCII into mojibake downstream)."""
    for bad, good in SMART.items():
        t = t.replace(bad, good)
    return t


def runs_text(paragraphs):
    parts = []
    for p in paragraphs:
        for r in p.runs:
            if r.text:
                parts.append(f'**{r.text}**' if r.font.bold else r.text)
        parts.append('\n')
    text = ascii_fold(''.join(parts)).strip()
    return re.sub(r'\*\*(\s*)\*\*', r'\1', text).strip()      # rejoin split bold runs


def deck_slides():
    out = []
    for slide in Presentation(str(DECK)).slides:
        texts = [runs_text(sh.text_frame.paragraphs) for sh in slide.shapes
                 if sh.has_text_frame and sh.text_frame.text.strip()]
        notes = ascii_fold(slide.notes_slide.notes_text_frame.text).strip() if slide.has_notes_slide else ''
        # A dark title/divider slide has no kicker, so its title is the FIRST
        # text box; every other kind starts with the kicker.
        dark = slide.background is not None and len(texts) <= 2
        title = texts[0] if dark else (texts[1] if len(texts) > 1 else (texts[0] if texts else ''))
        pictures = [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
        out.append({'texts': texts, 'notes': notes, 'title': title,
                    'subtitle': texts[1] if dark and len(texts) > 1 else '',
                    'pictures': pictures})
    return out


def norm(t):
    return re.sub(r'[^a-z0-9]+', ' ', (t or '').replace('*', '').lower()).strip()


def match_slides(cs, ds):
    """[(content, deck)] in DECK order, plus what could not be matched."""
    remaining = list(cs)
    pairs, unmatched = [], []
    for d in ds:
        hit = next((c for c in remaining if norm(c.get('title')) == norm(d['title'])), None)
        if hit is None:                       # title edited too: fall back to notes
            hit = next((c for c in remaining if c.get('notes') and d['notes']
                        and norm(c['notes'])[:60] == norm(d['notes'])[:60]), None)
        if hit is None:
            unmatched.append(d)
        else:
            remaining.remove(hit)
            pairs.append((hit, d))
    return pairs, unmatched, remaining


def body_texts(texts):
    """Drop kicker, title and the bold marker/letter labels."""
    def is_label(t):
        return re.fullmatch(r'[0-9A-Z]', t.replace('*', '').strip()) is not None
    return [t for t in texts[2:] if not is_label(t)]


def compare(c, d):
    diffs = []
    body = body_texts(d['texts'])
    kind = c['kind']
    if kind == 'steps':
        fields = [('steps', i, s) for i, s in enumerate(c['steps'])]
        if c.get('tip'):
            fields.append(('tip', None, c['tip']))
    elif kind == 'overview':
        fields = [('zones', i, z['label']) for i, z in enumerate(c['zones'])]
    elif kind in ('concept', 'placeholder'):
        fields = [('bullets', i, b) for i, b in enumerate(c.get('bullets', []))]
    else:
        fields = []                          # glossary, summary, title, divider: notes only
    if c.get('title') and norm(c['title']) != norm(d['title']):
        diffs.append(('title', c['title'], plain_title(d['title'])))
    if c.get('subtitle') and d.get('subtitle') and norm(c['subtitle']) != norm(d['subtitle']):
        diffs.append(('subtitle', c['subtitle'], d['subtitle']))
    for (name, idx, old), new in zip(fields, body):
        if old.strip() != new.strip():
            diffs.append((f'{name}[{idx}]' if idx is not None else name, old, new))
    if d['notes'] and d['notes'] != (c.get('notes') or '').strip():
        diffs.append(('notes', c.get('notes', ''), d['notes']))
    return diffs


def js_string(s):
    r"""A JS single-quoted literal. Line breaks are kept as \n -- the generator
    renders them as line breaks, and flattening them loses the author's layout."""
    body = s.replace('\\', '\\\\').replace("'", "\\'").replace('\n', '\\n')
    return "'" + body + "'"


def slide_blocks(src):
    """Source text of every slide entry, keyed by id, comments above kept with it."""
    arr = src.index('slides: [')
    first = src.index('    {\n', arr)
    blocks, order, i = {}, [], first
    while True:
        start = src.find('    {\n', i)
        if start == -1:
            break
        cstart = start
        while True:
            line_start = src.rfind('\n', 0, cstart - 1) + 1
            if src[line_start:cstart].lstrip().startswith('//'):
                cstart = line_start
            else:
                break
        depth, j = 0, start
        while j < len(src):
            if src[j] == '{':
                depth += 1
            elif src[j] == '}':
                depth -= 1
                if depth == 0:
                    break
            j += 1
        end = src.find('\n', j) + 1
        block = src[cstart:end]
        m = re.search(r"id: '([^']+)'", block)
        if not m:
            break
        blocks[m.group(1)] = block
        order.append(m.group(1))
        i = end
    head_end = min(blocks and src.index(blocks[order[0]]) or first, first)
    return blocks, order, src[:head_end], src[i:]


def plain_title(t):
    """Titles are rendered bold by the generator, so a title typed bold in
    PowerPoint comes back as '**Mounting a new Die**'. Strip the marks."""
    return re.sub(r'\*\*', '', t or '').strip()


def slugify(t):
    return re.sub(r'[^a-z0-9]+', '-', (t or 'slide').lower()).strip('-')[:40] or 'slide'


def import_new(d):
    """A slide added in PowerPoint: save its picture into shots/ and return a
    content block for it, so the next build keeps the slide instead of dropping it."""
    if not d['pictures']:
        print(f"  ! added slide {d['title']!r} has no picture -- add it to the content file by hand")
        return None, None
    title = plain_title(d['title'])
    slug = slugify(title)
    pic = d['pictures'][0]
    ext = pic.image.ext or 'png'
    name = f'manual_{slug}.{ext}'
    (SHOTS / name).write_bytes(pic.image.blob)
    body = body_texts(d['texts'])
    caption = body[0] if body else ''
    lines = [
        '    {',
        f"      id: {js_string(slug)}, kind: 'image', kicker: 'Team lead', title: {js_string(title)},",
        f'      image: {js_string(name)},',
    ]
    if caption:
        lines.append(f'      caption: {js_string(caption)},')
    lines.append(f"      notes: {js_string(d['notes'] or 'Added in PowerPoint, imported by pull_pptx_edits.py.')},")
    lines.append('    },')
    block = '\n'.join(lines) + '\n'
    print(f'  imported added slide {title!r} -> {name} (id {slug})')
    return slug, block


def apply_edits(all_diffs, deck_order, new_blocks):
    src = CONTENT.read_text(encoding='utf-8')
    for slide_id, diffs in all_diffs:
        start = src.index(f"id: '{slide_id}'")
        end = src.find("\n    {\n", start)
        end = len(src) if end == -1 else end
        block = src[start:end]
        for label, old, new in diffs:
            if label == 'notes':
                block = re.sub(r"notes: (?:'(?:[^'\\]|\\.)*'(?:\s*\+\s*)?)+,",
                               'notes: ' + js_string(new) + ',', block, count=1)
            elif label in ('title', 'subtitle'):
                block = re.sub(label + r": '(?:[^'\\]|\\.)*'", label + ': ' + js_string(new), block, count=1)
            else:
                needle = js_string(old)
                if needle in block:
                    block = block.replace(needle, js_string(new), 1)
                else:
                    print(f'  ! {slide_id} {label}: old text not found, left alone')
        src = src[:start] + block + src[end:]
    CONTENT.write_text(src, encoding='utf-8')

    blocks, order, head, tail = slide_blocks(CONTENT.read_text(encoding='utf-8'))
    blocks.update(new_blocks)
    if set(deck_order) == set(blocks) and deck_order != order:
        CONTENT.write_text(head + ''.join(blocks[i] for i in deck_order) + tail, encoding='utf-8')
        print('  content file reordered to match the deck')
    elif set(deck_order) != set(blocks):
        print('  ! order not applied: the deck and the content file hold different slides')


def main():
    cs, ds = content_slides(), deck_slides()
    pairs, unmatched, leftover = match_slides(cs, ds)
    for d in unmatched:
        print(f"! deck slide not in the content file (added in PowerPoint?): {d['title']!r}")
    for c in leftover:
        print(f"! content slide missing from the deck: {c['id']} ({c.get('title')!r})")
    order_now = [c['id'] for c in cs]
    deck_order = [c['id'] for c, _ in pairs]
    if deck_order != order_now:
        print(f'\nslide order changed:\n  was: {" ".join(order_now)}\n  now: {" ".join(deck_order)}')
    all_diffs = []
    for c, d in pairs:
        diffs = compare(c, d)
        if diffs:
            all_diffs.append((c['id'], diffs))
            print(f"\n{c['id']}")
            for label, old, new in diffs:
                print(f'  {label}\n    - {old}\n    + {new}')
    if not all_diffs and deck_order == order_now:
        print('no differences')
        return
    if '--apply' in sys.argv:
        new_blocks = {}
        for d in unmatched:
            slug, block = import_new(d)
            if slug:
                new_blocks[slug] = block
                deck_order.insert(ds.index(d), slug)
        apply_edits(all_diffs, deck_order, new_blocks)
        print(f'\napplied to {CONTENT.name}')
    else:
        print('\nrun again with --apply to write these into the content file')


main()
